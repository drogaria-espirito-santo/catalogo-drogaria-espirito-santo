from __future__ import annotations

import json
import re
import shutil
import unicodedata
from datetime import date
from pathlib import Path

import openpyxl
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"C:\Users\Acer\Downloads\Catalogo_para_filtrar_fotos.xlsx")
PHOTOS_SOURCE = Path(r"C:\Users\Acer\Desktop\Fotos drogaria")

DATA_DIR = ROOT / "src" / "data"
PUBLIC_PLACEHOLDERS = ROOT / "public" / "placeholders"
PUBLIC_PRODUCTS = ROOT / "public" / "produtos"

SUPPORTED_IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp"}

CATEGORY_ORDER = [
    "Mobilidade e Ortopedia",
    "Curativo e Enfermagem",
    "Incontinência e Cuidado ao Leito",
    "Aferição e Diagnóstico",
    "Respiratório",
    "Diabetes e Testes",
    "Termoterapia e Conforto",
    "Meias de Compressão",
]

BRAND_RED = "#ff1717"
BRAND_BLUE = "#1598cf"
BRAND_PINK = "#ed87ad"
INK = "#172026"
SLATE = "#64748b"


def slugify(value: object) -> str:
    text = str(value or "").strip().lower()
    text = unicodedata.normalize("NFD", text)
    text = "".join(char for char in text if unicodedata.category(char) != "Mn")
    text = re.sub(r"[^a-z0-9]+", "-", text)
    return text.strip("-") or "item"


def clean(value: object) -> str:
    text = str(value or "").strip()
    if "Ã" in text or "Â" in text:
        try:
            text = text.encode("latin1").decode("utf-8")
        except UnicodeError:
            pass
    return "" if text in {"—", "-", "None"} else text


def title_case_product(name: str) -> str:
    small = {"c", "d", "de", "do", "da", "das", "dos", "em", "p", "s"}
    out = []
    for token in name.split():
        raw = token.strip()
        base = raw.lower()
        if base in small or any(ch.isdigit() for ch in raw):
            out.append(raw)
        elif len(raw) <= 3 and raw.isalpha():
            out.append(raw)
        else:
            out.append(raw[:1].upper() + raw[1:].lower())
    return " ".join(out)


def load_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\segoeuib.ttf" if bold else r"C:\Windows\Fonts\segoeui.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def wrap_text(text: str, font: ImageFont.ImageFont, max_width: int) -> list[str]:
    words = text.replace("&", "e").split()
    lines: list[str] = []
    current = ""
    probe = ImageDraw.Draw(Image.new("RGB", (1, 1)))
    for word in words:
        next_line = f"{current} {word}".strip()
        bbox = probe.textbbox((0, 0), next_line, font=font)
        if bbox[2] <= max_width or not current:
            current = next_line
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines[:2]


def draw_cross(draw: ImageDraw.ImageDraw, cx: int, cy: int, size: int, color: str) -> None:
    w = size // 3
    draw.rounded_rectangle((cx - w // 2, cy - size // 2, cx + w // 2, cy + size // 2), radius=12, fill=color)
    draw.rounded_rectangle((cx - size // 2, cy - w // 2, cx + size // 2, cy + w // 2), radius=12, fill=color)


def draw_icon(draw: ImageDraw.ImageDraw, label: str) -> None:
    key = slugify(label)
    red = BRAND_RED
    blue = BRAND_BLUE
    pink = BRAND_PINK
    dark = "#263238"
    stroke = 24

    if "cadeira" in key or "rodas" in key:
        draw.ellipse((210, 415, 365, 570), outline=blue, width=stroke)
        draw.ellipse((535, 460, 650, 575), outline=pink, width=stroke)
        draw.line((335, 420, 470, 420, 545, 535), fill=dark, width=stroke)
        draw.line((455, 315, 455, 420), fill=dark, width=stroke)
        draw.line((455, 330, 585, 330), fill=dark, width=stroke)
        draw.line((585, 330, 620, 450), fill=dark, width=stroke)
        return

    if "andador" in key:
        draw.line((255, 235, 255, 585), fill=dark, width=stroke)
        draw.line((645, 235, 645, 585), fill=dark, width=stroke)
        draw.line((255, 235, 645, 235), fill=blue, width=stroke)
        draw.line((285, 390, 615, 390), fill=red, width=stroke)
        draw.line((255, 585, 190, 665), fill=dark, width=stroke)
        draw.line((645, 585, 710, 665), fill=dark, width=stroke)
        return

    if "bengala" in key or "muleta" in key:
        draw.arc((335, 175, 555, 365), 180, 360, fill=blue, width=stroke)
        draw.line((445, 270, 445, 660), fill=dark, width=stroke)
        draw.line((385, 515, 505, 515), fill=red, width=stroke)
        draw.line((445, 660, 395, 710), fill=dark, width=stroke)
        return

    if "joelheira" in key or "munhequeira" in key or "cotoveleira" in key or "tornozeleira" in key:
        draw.rounded_rectangle((275, 190, 625, 660), radius=64, fill="#20242a")
        draw.rounded_rectangle((315, 230, 585, 620), radius=48, outline=pink, width=stroke)
        draw.ellipse((385, 370, 515, 500), fill="#ffffff")
        draw.rounded_rectangle((360, 540, 540, 590), radius=16, fill=blue)
        return

    if "curativo" in key or "gaze" in key or "atadura" in key or "esparadrapo" in key:
        draw.rounded_rectangle((190, 350, 710, 535), radius=72, fill="#fff1f2", outline=red, width=stroke)
        draw.rounded_rectangle((390, 315, 510, 570), radius=36, fill="#ffffff", outline=blue, width=18)
        draw_cross(draw, 450, 442, 88, red)
        return

    if "nebul" in key or "respir" in key or "umid" in key or "oxigen" in key:
        draw.ellipse((240, 240, 450, 510), fill="#e0f2fe", outline=blue, width=stroke)
        draw.ellipse((450, 240, 660, 510), fill="#e0f2fe", outline=blue, width=stroke)
        draw.line((450, 395, 450, 630), fill=dark, width=stroke)
        draw.rounded_rectangle((330, 610, 570, 700), radius=35, fill=pink)
        return

    if "glic" in key or "diabetes" in key or "lanceta" in key or "tira" in key or "insulina" in key:
        draw.rounded_rectangle((250, 210, 430, 650), radius=42, fill="#ffffff", outline=blue, width=stroke)
        draw.rounded_rectangle((295, 285, 385, 370), radius=16, fill="#dbeafe")
        draw.line((520, 220, 625, 450), fill=dark, width=stroke)
        draw.ellipse((500, 455, 660, 615), fill="#ffe4e6", outline=red, width=stroke)
        draw_cross(draw, 580, 535, 70, red)
        return

    if "pressao" in key or "termometro" in key or "oximetro" in key or "estetoscopio" in key:
        draw.rounded_rectangle((225, 235, 675, 560), radius=54, fill="#ffffff", outline=blue, width=stroke)
        draw.line((305, 420, 405, 420, 455, 340, 525, 500, 595, 420), fill=red, width=18)
        draw.rounded_rectangle((350, 600, 550, 685), radius=28, fill=pink)
        return

    if "fralda" in key or "leito" in key or "comadre" in key or "almofada" in key or "coxim" in key:
        draw.rounded_rectangle((200, 380, 700, 570), radius=42, fill="#ffffff", outline=blue, width=stroke)
        draw.rounded_rectangle((220, 315, 390, 400), radius=28, fill="#ffe4e6", outline=pink, width=16)
        draw.line((230, 570, 230, 650), fill=dark, width=stroke)
        draw.line((670, 570, 670, 650), fill=dark, width=stroke)
        return

    if "meia" in key:
        draw.line((350, 185, 350, 520), fill=dark, width=90)
        draw.line((395, 520, 610, 520), fill=dark, width=90)
        draw.arc((515, 405, 700, 590), 0, 95, fill=dark, width=90)
        draw.line((315, 250, 385, 250), fill=pink, width=18)
        draw.line((315, 315, 385, 315), fill=blue, width=18)
        return

    draw.rounded_rectangle((245, 235, 655, 615), radius=58, fill="#ffffff", outline=blue, width=stroke)
    draw_cross(draw, 450, 425, 170, red)
    draw.rounded_rectangle((315, 615, 585, 690), radius=28, fill=pink)


def generate_placeholder_png(label: str, output: Path) -> None:
    image = Image.new("RGB", (900, 900), "#ffffff")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 900, 150), fill=BRAND_RED)
    draw.rounded_rectangle((44, 190, 856, 812), radius=52, fill="#f8fafc", outline="#e2e8f0", width=4)
    draw.ellipse((645, 205, 845, 405), fill="#e0f2fe")
    draw.ellipse((52, 590, 245, 783), fill="#ffe4e6")
    draw_icon(draw, label)

    title_font = load_font(54, bold=True)
    caption_font = load_font(30, bold=True)
    lines = wrap_text(label, title_font, 690)
    y = 700
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=title_font)
        draw.text(((900 - (bbox[2] - bbox[0])) / 2, y), line, fill=INK, font=title_font)
        y += 60
    caption = "Imagem padrão da categoria"
    bbox = draw.textbbox((0, 0), caption, font=caption_font)
    draw.text(((900 - (bbox[2] - bbox[0])) / 2, 800), caption, fill=SLATE, font=caption_font)
    image.save(output, "PNG", optimize=True)


def prepare_product_photos() -> dict[str, str]:
    PUBLIC_PRODUCTS.mkdir(parents=True, exist_ok=True)
    if not PHOTOS_SOURCE.exists():
        return {}

    matches: dict[str, str] = {}
    for source in PHOTOS_SOURCE.iterdir():
        if not source.is_file() or source.suffix.lower() not in SUPPORTED_IMAGE_EXTENSIONS:
            continue
        key = slugify(source.stem)
        destination = PUBLIC_PRODUCTS / f"{key}{source.suffix.lower()}"
        shutil.copy2(source, destination)
        matches[key] = f"/produtos/{destination.name}"
    return matches


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(f"Planilha não encontrada: {SOURCE}")

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    PUBLIC_PLACEHOLDERS.mkdir(parents=True, exist_ok=True)

    for stale in [*PUBLIC_PLACEHOLDERS.glob("*.svg"), *PUBLIC_PLACEHOLDERS.glob("*.png")]:
        stale.unlink()

    product_photos = prepare_product_photos()

    wb = openpyxl.load_workbook(SOURCE, data_only=True)
    ws = wb["Catálogo"]
    rows = list(ws.iter_rows(values_only=True))
    header = [clean(value) for value in rows[0]]
    products = []
    seen_codes = set()
    subtypes = []
    real_photo_count = 0

    for row in rows[1:]:
        item = dict(zip(header, row))
        categoria = clean(item.get("Categoria"))
        subtipo = clean(item.get("Subtipo"))
        codigo = clean(item.get("Códigos Internos"))
        produto_raw = clean(item.get("Produto"))
        if not categoria or not subtipo or not codigo or not produto_raw:
            continue

        codigo = re.sub(r"\.0$", "", codigo)
        base_code = codigo
        suffix = 2
        while codigo in seen_codes:
            codigo = f"{base_code}-{suffix}"
            suffix += 1
        seen_codes.add(codigo)

        subtype_slug = slugify(subtipo)
        if subtipo not in subtypes:
            subtypes.append(subtipo)

        product_key = slugify(produto_raw)
        image_path = product_photos.get(product_key, f"/placeholders/{subtype_slug}.png")
        if image_path.startswith("/produtos/"):
            real_photo_count += 1

        product = {
            "id": f"{codigo}-{product_key}",
            "codigo": codigo,
            "produto": title_case_product(produto_raw),
            "produtoOriginal": produto_raw,
            "categoria": categoria,
            "categoriaSlug": slugify(categoria),
            "subtipo": subtipo,
            "subtipoSlug": subtype_slug,
            "variante": clean(item.get("Variante")),
            "temFotoHoje": clean(item.get("Tem foto hoje?")).lower() == "sim",
            "imagem": image_path,
            "whatsapp": "5527995050105",
        }
        products.append(product)

    categories = []
    for category in CATEGORY_ORDER:
        category_products = [p for p in products if p["categoria"] == category]
        if not category_products:
            continue
        category_subtypes = sorted({p["subtipo"] for p in category_products})
        categories.append(
            {
                "nome": category,
                "slug": slugify(category),
                "total": len(category_products),
                "subtipos": [
                    {
                        "nome": subtype,
                        "slug": slugify(subtype),
                        "total": sum(1 for p in category_products if p["subtipo"] == subtype),
                        "imagem": f"/placeholders/{slugify(subtype)}.png",
                    }
                    for subtype in category_subtypes
                ],
                "imagem": f"/placeholders/{slugify(category_subtypes[0])}.png",
            }
        )

    for subtype in subtypes:
        generate_placeholder_png(subtype, PUBLIC_PLACEHOLDERS / f"{slugify(subtype)}.png")

    payload = {
        "geradoEm": date.today().isoformat(),
        "total": len(products),
        "fotosReais": real_photo_count,
        "categorias": categories,
        "produtos": products,
    }
    (DATA_DIR / "produtos.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(
        f"Importados {len(products)} produtos em {len(categories)} categorias. "
        f"Fotos reais vinculadas: {real_photo_count}."
    )


if __name__ == "__main__":
    main()
